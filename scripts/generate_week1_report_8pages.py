from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "complete-week-1-report-8-pages.docx"
FIG1 = DOCS / "generated-assets" / "figure-runtime-topology.png"
FIG2 = DOCS / "generated-assets" / "figure-daily-flow.png"


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_end)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    add_page_number(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11.2)
    pf = normal.paragraph_format
    pf.space_after = Pt(5)
    pf.line_spacing = 1.05

    for style_name, size in [("Heading 1", 15.5), ("Heading 2", 12.8), ("Heading 3", 11.3)]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(7)
        style.paragraph_format.space_after = Pt(4)

    for style_name in ["List Bullet", "List Number"]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11.0)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.03


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("\n")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("OpenClaw Content Sentinel")
    r.bold = True
    r.font.size = Pt(21)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Week 1 Research Report and Architecture Plan")
    r.bold = True
    r.font.size = Pt(14)

    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    headers = ["Project Goal", "Constraint", "Primary Deliverable"]
    values = [
        "Autonomous daily competitor-aware content agent built on OpenClaw",
        "Zero additional cost beyond the LLM",
        "Professional 8-page architecture report with Week 1 plan",
    ]
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    for idx, text in enumerate(values):
        table.rows[1].cells[idx].text = text

    for line in [
        "Course Context: Intern Project Assignment",
        "Deliverable Type: Professional Week 1 Architecture Report",
        "Prepared on: March 25, 2026",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(line)

    doc.add_page_break()


def add_paragraph(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)


def main() -> None:
    doc = Document()
    style_document(doc)
    add_title(doc)

    doc.add_heading("1. Executive Summary", level=1)
    add_paragraph(
        doc,
        "OpenClaw Content Sentinel is a feasible Week 1 initiative if the system is designed as an OpenClaw-first control plane rather than a loose collection of scripts. The right MVP keeps OpenClaw responsible for the gateway runtime, cron and heartbeat scheduling, browser control, Telegram interaction, workspace memory, and skill loading. That choice reduces architectural sprawl, keeps the project aligned with the assignment, and makes later hardening easier."
    )
    add_paragraph(
        doc,
        "The strongest design constraint is the zero-cost requirement. That removes paid social APIs, paid search APIs, hosted orchestration services, and premium scraping layers. The stack therefore has to rely on public RSS feeds, public search pages, a self-hosted SearXNG-style metasearch layer, browser automation through persistent profiles, local storage, and the LLM selected by the company. This is realistic for research, drafting, and operator review. It is technically possible for posting too, but posting carries the highest operational and policy risk and must be treated as a controlled capability rather than an assumed entitlement."
    )
    add_bullets(
        doc,
        [
            "Pure OpenClaw is sufficient for the Week 1 MVP.",
            "LangGraph and local vector retrieval are valuable extensions, not day-one dependencies.",
            "Telegram should act as the operator console for approval, overrides, logs, and health checks.",
            "Human approval should remain mandatory before publication on LinkedIn, Facebook, and X.",
        ],
    )

    doc.add_heading("2. OpenClaw-First Architecture", level=1)
    add_paragraph(
        doc,
        "OpenClaw already provides most of the primitives the project brief asks for. The gateway owns channel integrations, tool routing, sessions, browser management, scheduler state, and the local agent workspace. Skills are loaded from bundled, managed, and workspace directories with workspace skills taking precedence, which makes custom `SKILL.md` packages the correct way to add competitor ingestion, trend lookup, content generation, approval control, and browser posting."
    )
    add_paragraph(
        doc,
        "The browser layer is especially important because the assignment requires ethical competitor review and browser-based posting while staying inside OpenClaw. Current OpenClaw documentation supports dedicated Chromium-family profiles controlled through the gateway. That gives the project a clean operational split: operator logins happen once in a dedicated profile, while the agent later reuses that profile for deterministic low-volume actions. The same browser surface can inspect public competitor pages, collect screenshots, and confirm whether a publish action actually completed."
    )
    add_paragraph(
        doc,
        "Scheduling should use both native mechanisms rather than forcing everything into a single loop. Cron is the correct trigger for the daily content run because it persists jobs and wakes the agent at exact times. Heartbeat is better for low-cost periodic checks such as failed-post detection, approval reminders, stale draft alerts, and auth drift detection. Telegram is then the human-facing layer that exposes `/status`, `/run-now`, `/approve`, `/reject`, `/logs`, and `/pause` without requiring server access."
    )
    doc.add_heading("2.1 Telegram Operator Model", level=2)
    add_paragraph(
        doc,
        "Telegram should be framed as the operational cockpit for the whole system, not just as a notification sink. The bot should expose a small but complete command set covering health, approval, manual triggering, and operational overrides. The most important commands for Week 1 are `/status`, `/run-now`, `/draft`, `/approve`, `/reject`, `/logs`, `/pause`, and `/resume`. Inline approval buttons should mirror the same actions so the operator can approve from mobile without typing commands manually."
    )
    add_bullets(
        doc,
        [
            "Use strict pairing or allowlist rules so only authorized operator IDs can interact with the bot.",
            "Send a draft preview with the source URL, selected angle, confidence score, and posting targets before approval.",
            "Return compact but actionable failure summaries instead of raw stack traces in chat messages.",
        ],
    )
    add_paragraph(
        doc,
        "This operator pattern is one of the best arguments for using OpenClaw natively. The same runtime that triggers the job can also collect the result, present it through Telegram, and persist the context in the workspace. That reduces glue code and makes the system easier to audit."
    )
    if FIG1.exists():
        doc.add_picture(str(FIG1), width=Inches(6.85))
        add_caption(doc, "Figure 1. Runtime topology showing the OpenClaw gateway as the control center.")
    add_paragraph(
        doc,
        "Figure 1 highlights the key architectural decision: OpenClaw is not just hosting prompts, it is acting as the operating system for the agent. That matters because the report must show an OpenClaw-first design, not a Python-first or Playwright-first design with OpenClaw added on top."
    )

    doc.add_heading("3. Zero-Cost Integration Strategy", level=1)
    add_paragraph(
        doc,
        "The zero-cost strategy has to be explicit and auditable. For competitor ingestion, the operator supplies a public article URL and the agent opens it through the OpenClaw browser, extracts the rendered article body, and stores a clean summary with source metadata such as title, date, URL, key claims, and outbound links. Public content only should be in scope for Week 1. No paywalls, CAPTCHA bypasses, or hidden content extraction should be attempted."
    )
    add_paragraph(
        doc,
        "Trending research should use multiple weak but free signals rather than pretending one source is authoritative. The best baseline is a combination of industry RSS feeds, public news searches, site-specific search queries, and a self-hosted SearXNG-style search layer for broad discovery. Google Trends scraping can be used as a weak signal, but it should not control the editorial angle alone because scraping-based trend tools are brittle and rate-limited. The now-archived `pytrends` project is a reminder that this part of the stack should stay optional."
    )
    add_paragraph(
        doc,
        "Content generation should follow a deterministic pipeline: extract source, gather external signals, compare claims, identify gaps or alternative angles, produce a long-form draft, then derive platform-specific rewrites for LinkedIn, Facebook, and X. Week 1 should also introduce a confidence score before any preview is sent. High-confidence drafts move to Telegram approval, medium-confidence drafts are flagged for review, and low-confidence runs stop with an exception summary instead of producing risky public copy."
    )
    add_bullets(
        doc,
        [
            "Image generation should stay local, using an open model through a local runner such as ComfyUI or a similar self-hosted setup.",
            "If local image generation is not ready in Week 1, a branded text-card fallback is preferable to a paid image API.",
            "Every run should log which research sources were used so the team can prove zero non-LLM API cost.",
        ],
    )
    add_paragraph(
        doc,
        "Posting is the most sensitive zero-cost capability because the only practical path is browser automation with persistent sessions. The correct operating model is one dedicated browser profile per platform, manual first login by the operator, low posting frequency, screenshot verification before and after submit, and immediate abort if the site presents a checkpoint, challenge page, or suspicious-activity warning."
    )
    doc.add_heading("3.1 Browser Posting and Anti-Detection Practice", level=2)
    add_paragraph(
        doc,
        "The anti-detection goal is not to imitate human randomness. The correct goal is to behave like a stable, low-volume operator workflow. The browser should use one persistent profile per platform, one posting action at a time, deterministic page-state checks, conservative retries, and immediate escalation when the platform asks for re-authentication or presents a trust challenge. Community Playwright skills can help with waiting patterns and selector recovery, but the final posting logic should remain narrow and intentionally boring."
    )
    add_bullets(
        doc,
        [
            "Manual first login, then persistent session reuse.",
            "One tab and one submit path per platform to limit state drift.",
            "Screenshot before submit, screenshot after confirmation, and structured result logging.",
            "Abort on challenge pages, security interstitials, or unusual compose-form changes.",
        ],
    )
    add_paragraph(
        doc,
        "This is also where the zero-cost claim needs evidence. The run log should show that the system used native browser automation and public sources rather than paid APIs. That record will matter during the final 30-day proof of zero marginal cost."
    )

    doc.add_heading("4. Community Skills: Adapt vs Rebuild", level=1)
    add_paragraph(
        doc,
        "The brief mentions community references such as `linkedin-automation`, `meta-fb-inbox`, `ai-hunter-pro`, and `playwright-scraper`. Those are useful references, but not trustworthy production dependencies for this system. The project should adapt patterns from them and rebuild the business-critical parts in a narrow local codebase."
    )
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for idx, text in enumerate(["Area", "Adapt", "Rebuild"]):
        cell = table.rows[0].cells[idx]
        cell.text = text
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    rows = [
        ("Browser control", "Selector retries, wait logic, screenshot discipline", "Final posting flows and account handling"),
        ("Scraping", "Rendered page extraction patterns", "Competitor-specific cleaning and prompt-injection filtering"),
        ("Operations", "Result payload shapes and approval ergonomics", "Logging, confidence gating, retries, and escalation"),
        ("Security", "Nothing security-critical should be copied blindly", "Credential handling, session isolation, and permission limits"),
    ]
    for area, adapt, rebuild in rows:
        cells = table.add_row().cells
        cells[0].text = area
        cells[1].text = adapt
        cells[2].text = rebuild
    add_paragraph(
        doc,
        "This split is important for maintainability. Community skills can show how others approached persistence and browser ergonomics, but they do not remove the need for a clean project-owned implementation that can be audited, fixed quickly, and defended in a production review."
    )

    doc.add_page_break()
    doc.add_heading("5. Pure OpenClaw vs Hybrid LangGraph Stack", level=1)
    add_paragraph(
        doc,
        "Pure OpenClaw is the right Week 1 choice because it already covers the gateway, browser, scheduling, Telegram integration, workspace memory, and the skill model. That keeps the initial system legible. A hybrid OpenClaw plus LangGraph stack becomes attractive only once the workflow needs durable checkpoints, resumable state, or explicit multi-agent orchestration such as Researcher, Analyzer, Writer, Editor, and Publisher. LangGraph is therefore best treated as a Week 2 optimization for control and observability, not a prerequisite for demonstrating architecture maturity."
    )
    add_paragraph(
        doc,
        "The same logic applies to retrieval. Native OpenClaw memory is enough for operator notes, recent runs, approved phrasing, and short historical context. A local vector store such as Chroma or FAISS becomes useful once the team wants retrieval over a larger archive of competitor posts and historical outputs. A graph store is even further from the Week 1 critical path. It can be useful for long-term topic relationships, but it does not improve the first deployment enough to justify the extra operational load."
    )
    add_paragraph(
        doc,
        "The recommended architectural decision is therefore incrementalism. Week 1 proves the single-agent OpenClaw loop. Week 2 introduces structured orchestration only if the approval flow, retries, or drafting stages begin to feel opaque. This order matters because it prevents the team from spending the first week wiring abstractions instead of proving that the basic content loop works."
    )

    doc.add_heading("6. Security, Reliability, and Compliance", level=1)
    add_paragraph(
        doc,
        "Security in this project is mostly about limiting trust. Competitor pages, search results, and community skills should all be treated as untrusted input. The recommended approach is to run scraping, parsing, and text processing in sandboxed skills, keep write permissions minimal, separate extraction from generation, and prevent scraped text from influencing tool-choice authority. Browser sessions should live in dedicated OpenClaw-managed profiles and raw credentials should stay outside the workspace in configuration or environment storage."
    )
    add_paragraph(
        doc,
        "Reliability depends on explicit controls rather than model optimism. Every run should record a run ID, trigger source, source URL, research inputs, chosen angle, model used, confidence score, approval status, post result per platform, screenshot paths, and retry counts. Telegram should surface exceptions in plain language so the operator can act without opening server logs. The project should also expose a draft-only fallback mode, a platform-level circuit breaker, and a manual re-run command to prevent repeated failures from escalating into account problems."
    )
    doc.add_heading("6.1 Prompt Injection and Data Hygiene", level=2)
    add_paragraph(
        doc,
        "Prompt injection risk is especially relevant here because competitor articles and public search pages are both model-facing inputs. The safe pattern is to treat scraped content as data, never as instruction. The extraction step should clean navigation noise, isolate article text, remove or neutralize imperative phrasing that tries to influence tools, and pass a structured summary to the drafting prompt. The model can analyze the content, but it should not be allowed to inherit authority from the webpage."
    )
    add_paragraph(
        doc,
        "Data hygiene also affects attribution and originality. The system should preserve source metadata, list the supporting URLs used in trend research, and avoid reproducing large passages of competitor content. The goal is commentary and synthesis, not hidden rewriting. That distinction is central to both ethics and brand safety."
    )
    if FIG2.exists():
        doc.add_picture(str(FIG2), width=Inches(6.75))
        add_caption(doc, "Figure 2. Controlled daily execution with a mandatory approval gate before posting.")
    add_paragraph(
        doc,
        "The legal and Terms-of-Service position has to be stated clearly in the report. Public article analysis is easier to justify when the system attributes sources and produces original commentary instead of copying text. Social posting through browser automation is more fragile because major platforms commonly restrict bots and unauthorized automated behavior. The honest production stance is therefore: technically feasible, operationally fragile, policy-sensitive, and only acceptable with explicit stakeholder approval, low-volume operation, and a human-in-the-loop posting gate."
    )
    add_bullets(
        doc,
        [
            "Mandatory controls: approval before publish, low-frequency posting, screenshot evidence, and secure profile storage.",
            "Operational controls: retries with backoff, challenge-page detection, pause and resume commands, and exception alerts.",
            "Compliance controls: attribution, disclosure where appropriate, and no hidden scraping or credential leakage.",
        ],
    )

    doc.add_page_break()
    doc.add_heading("7. Implementation Roadmap and Week 1 Deliverables", level=1)
    add_paragraph(
        doc,
        "Week 1 should establish the skeleton that makes the rest of the assignment realistic. That means deploying OpenClaw on a persistent host, connecting Telegram, creating the workspace, defining the command surface, and scaffolding the first skills for URL ingestion, trend lookup, content drafting, and approval control. The target is not full autonomy in Week 1; the target is a stable architecture that can run a supervised end-to-end dry run."
    )
    doc.add_heading("7.1 Base Deployment Blueprint", level=2)
    add_paragraph(
        doc,
        "The preferred host for Week 1 is a persistent Ubuntu VM or server with Docker available for sandboxed skill execution and enough storage for browser profiles, logs, screenshots, and generated artifacts. The deployment should create a private workspace repository, define the OpenClaw profile, configure the selected LLM credentials, enable Telegram, and verify that cron jobs survive service restarts. This gives the team a production-like baseline from the start instead of building first on a transient laptop session."
    )
    add_bullets(
        doc,
        [
            "Host OpenClaw on a persistent machine with browser support and secure local storage.",
            "Keep secrets out of the workspace and store them in environment or OpenClaw config.",
            "Initialize the skills directory early so each Week 1 capability maps to a concrete `SKILL.md` package.",
            "Validate cron, heartbeat, browser profile creation, and Telegram routing before building content logic.",
        ],
    )
    roadmap = doc.add_table(rows=1, cols=3)
    roadmap.style = "Table Grid"
    for idx, text in enumerate(["Phase", "Main Outcome", "Output"]):
        cell = roadmap.rows[0].cells[idx]
        cell.text = text
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    roadmap_rows = [
        ("Week 1", "Deployment, Telegram, base skills, research report", "Dry run to Telegram preview"),
        ("Week 2", "Content pipeline, optional LangGraph, browser posting skills", "Platform-specific post preparation"),
        ("Week 3", "Daily scheduler, confidence gates, recovery and approvals", "Reliable automation loop"),
        ("Week 4", "Simulation, documentation, handover, Docker-compose package", "Production handoff set"),
    ]
    for phase, outcome, output in roadmap_rows:
        row = roadmap.add_row().cells
        row[0].text = phase
        row[1].text = outcome
        row[2].text = output

    add_paragraph(
        doc,
        "By the end of Day 2, the sync should resolve six decisions: deployment host, LLM provider, operator approval policy, local image-generation path, acceptance of platform automation risk, and whether LangGraph enters in Week 2 or stays deferred. Those decisions remove most of the ambiguity from the implementation backlog."
    )
    add_numbered(
        doc,
        [
            "Deploy and validate OpenClaw with workspace, model, cron, and heartbeat.",
            "Connect Telegram with strict access control and the initial command set.",
            "Implement competitor URL fetch and trend lookup prototypes.",
            "Produce a supervised draft and send it to Telegram for approval review.",
        ],
    )
    add_paragraph(
        doc,
        "A successful Week 1 close-out should therefore produce four visible artifacts: the architecture report, a working OpenClaw deployment, a Telegram-connected operator interface, and a dry-run flow that reaches a draft preview. Those deliverables are enough to justify moving into browser posting and deeper workflow orchestration in Week 2."
    )

    doc.add_page_break()
    doc.add_heading("8. Final Recommendation", level=1)
    add_paragraph(
        doc,
        "The assignment should move forward as a pure OpenClaw MVP with disciplined boundaries. OpenClaw should remain the orchestration layer, Telegram the operator console, local browser profiles the only posting mechanism, and free public web sources the only research inputs besides the chosen LLM. LangGraph, vector retrieval, and graph storage are useful enhancements, but the Week 1 report should recommend them as controlled later additions. The immediate objective is a reliable, auditable, zero-cost architecture that can produce daily drafts, expose human approval, and evolve into full automation without losing maintainability."
    )
    add_paragraph(
        doc,
        "Primary sources reviewed for this summary include the official OpenClaw documentation for skills, browser control, cron, Telegram, workspace architecture, and security, together with official or public policy references for LinkedIn, Meta/Facebook, and X. The report intentionally uses conservative wording where platform policy language is broad or subject to change."
    )

    doc.save(str(OUT))
    print(f"Created {OUT}")


if __name__ == "__main__":
    main()
